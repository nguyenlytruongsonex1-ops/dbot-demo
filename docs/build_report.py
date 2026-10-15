"""Tạo báo cáo docs/bao-cao-dbot-demo.docx bằng python-docx.
Chạy: python docs/build_report.py"""
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = "https://github.com/nguyenlytruongsonex1-ops/dbot-demo/tree/feature/realtime-streamlit-pg"

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin, s.right_margin = Cm(3), Cm(2)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(13)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.3
for name, size in (("Title", 18), ("Heading 1", 15), ("Heading 2", 13)):
    h = doc.styles[name]; h.font.name = "Times New Roman"; h.font.size = Pt(size)
    h.font.bold = name != "Title"; h.font.color.rgb = None
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    h.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

def p(text, align=None, italic=False):
    para = doc.add_paragraph()
    r = para.add_run(text); r.italic = italic
    if align: para.alignment = align
    else: para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return para

def bullet(text):
    doc.add_paragraph(text, style="List Bullet")

def caption(text):
    p(text, WD_ALIGN_PARAGRAPH.CENTER, italic=True)

def table(header, rows, cap, widths=None):
    caption(cap)
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
    for row in t.rows:
        for i, c in enumerate(row.cells):
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs: run.font.size = Pt(11)
            if widths: c.width = Cm(widths[i])
    doc.add_paragraph()

H = doc.add_heading
p("Mã nguồn: " + REPO, WD_ALIGN_PARAGRAPH.LEFT)
doc.add_paragraph("Báo cáo tổng hợp: Demo chẩn đoán và tối ưu cơ sở dữ liệu PostgreSQL bằng mô hình ngôn ngữ lớn, có giám sát thời gian thực", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
p("Tài liệu học tập. Thực hiện: Nguyễn Lý Trường Sơn. Ngày 15/10/2026.", WD_ALIGN_PARAGRAPH.CENTER)
p("Bản này tích hợp dự án \"LLM As DBA\" của Nhóm 3 (môn Hệ cơ sở dữ liệu tiên tiến). Luồng trình diễn ba kịch bản, cách hỏi AI theo hai loại prompt, thẻ kinh nghiệm, cách chấm điểm tự động và phần Mini λ-Tune được chuyển từ dự án đó sang PostgreSQL.", WD_ALIGN_PARAGRAPH.CENTER, italic=True)

H("1. Giới thiệu và mục tiêu", 1)
p("Quản trị cơ sở dữ liệu cần nhiều kinh nghiệm. Khi có sự cố, người quản trị (DBA) phải đọc nhiều số liệu, đoán nguyên nhân và chọn cách xử lý. Khi hệ thống chạy bình thường, DBA lại phải chỉnh nhiều tham số cấu hình để hệ thống nhanh hơn. Hai bài báo D-Bot và λ-Tune đề xuất dùng mô hình ngôn ngữ lớn (LLM) để hỗ trợ hai việc này.")
p("Báo cáo trình bày một bản demo để học, dựng lại ý tưởng của hai bài báo ở mức nhẹ. Đây không phải công trình nghiên cứu và không nhằm so sánh với kết quả của các bài báo. Bản demo có bốn mục tiêu. Thứ nhất, gây ra sự cố thật trên PostgreSQL bằng các file SQL kịch bản. Thứ hai, trình diễn từng bước cách AI chẩn đoán khi có và khi không có số đo, theo cách làm của Nhóm 3. Thứ ba, tự phát hiện sự cố và tự gọi AI phân tích theo thời gian thực. Thứ tư, chạy được trên máy cá nhân bằng Docker Compose.")

H("2. Tóm tắt các công trình liên quan", 1)
H("2.1. D-Bot", 2)
p("Bài báo D-Bot: Database Diagnosis System using Large Language Models (Zhou và cộng sự, arXiv 2312.01454) xây dựng hệ thống chẩn đoán sự cố cho PostgreSQL. Hệ thống trích xuất tri thức chẩn đoán từ tài liệu thành các thẻ kinh nghiệm, cung cấp công cụ cho LLM gọi để lấy số liệu, suy luận theo cây để thử nhiều hướng, và cho nhiều agent chuyên môn phối hợp. Bài báo dùng GPT-4 và hệ thống giám sát Prometheus, nên khá nặng để dựng lại trên máy cá nhân.")
H("2.2. λ-Tune", 2)
p("λ-Tune (Giannakouris và Trummer) dùng LLM để đề xuất cấu hình cơ sở dữ liệu, gồm tham số bộ nhớ và index, dựa trên mô tả workload và phần cứng. Hệ thống chạy thử từng cấu hình trên workload thật, dừng sớm cấu hình chạy chậm, rồi giữ cấu hình nhanh nhất. Ý chính là không tin ngay lời AI mà phải đo thật.")
H("2.3. Dự án LLM As DBA của Nhóm 3", 2)
p("Dự án của Nhóm 3 là một ứng dụng Streamlit chạy trên MySQL 8, dữ liệu là bảng nhân viên 1 triệu dòng. Ứng dụng có ba kịch bản: thiếu index, truy vấn con tương quan và tranh chấp khóa. Mỗi kịch bản đi theo trình tự gây sự cố, đo, hỏi AI theo hai cách, sửa rồi đo lại. Prompt A chỉ kể triệu chứng. Prompt B làm theo kiểu D-Bot, có thẻ kinh nghiệm và kết quả công cụ thật. Câu trả lời được chấm tự động bằng từ khóa. Ứng dụng còn có tab Mini λ-Tune. Dự án chưa có phần giám sát chạy nền và chưa lưu lịch sử.")

H("3. Kiến trúc hệ thống", 1)
p("Hệ thống gồm sáu dịch vụ trong Docker Compose. LLM có thể là Claude, Gemini, OpenAI qua API, hoặc Ollama chạy trên máy người dùng. Hình 1 mô tả các thành phần và luồng dữ liệu.")
doc.add_picture(os.path.join(HERE, "kien-truc.png"), width=Cm(16))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
caption("Hình 1. Kiến trúc tổng thể của bản demo")
table(["Dịch vụ", "Vai trò"], [
    ["postgres", "PostgreSQL 16, bật pg_stat_statements, chứa dữ liệu TPC-H"],
    ["tpch-seed", "Sinh dữ liệu TPC-H SF 0.1 bằng dbgen rồi nạp vào cơ sở dữ liệu, chỉ chạy một lần"],
    ["monitor", "Quét các view thống kê mỗi 5 giây, tạo alert và tự đánh dấu đã giải quyết"],
    ["backend", "API nội bộ viết bằng FastAPI: đọc alert, lưu cài đặt, chạy và dừng kịch bản, tải kịch bản lên"],
    ["dbot-worker", "Tự nhận alert mới, chẩn đoán bằng CrewAI với hai agent, lưu báo cáo"],
    ["streamlit", "Giao diện duy nhất ở cổng 8501, gồm bảy tab"],
], "Bảng 1. Các dịch vụ trong Docker Compose", [3.5, 11.5])
p("Alert, báo cáo và cài đặt được lưu trong SQLite trên một volume dùng chung. Vì vậy khi tải lại trang hoặc khởi động lại, lịch sử vẫn còn. So với bản demo trước, trang HTML riêng và kênh WebSocket được bỏ. Streamlit tự làm mới phần giám sát mỗi 3 giây bằng cơ chế fragment.")

H("4. Dữ liệu và kịch bản sự cố", 1)
H("4.1. Dữ liệu TPC-H", 2)
p("Dữ liệu là bộ TPC-H với scale factor 0.1, khoảng 100 MB, gồm 8 bảng. Bảng lớn nhất là lineitem với khoảng 600 nghìn dòng, tiếp theo là orders với 150 nghìn dòng. Các bảng chỉ có khóa chính. Tác giả cố ý không tạo index phụ trên các cột như l_partkey hay o_custkey để kịch bản thiếu index và truy vấn con tương quan có tác dụng. Dữ liệu được sinh khi cài đặt nên không lưu trong repo.")
H("4.2. Chuyển kịch bản từ MySQL sang PostgreSQL", 2)
p("Ba kịch bản của Nhóm 3 được viết lại trên bảng TPC-H. Bảng 2 so sánh cách làm ở hai hệ quản trị.")
table(["Kịch bản", "Bản gốc MySQL", "Bản PostgreSQL"], [
    ["Thiếu index", "Lọc emp theo dno, đo bằng EXPLAIN và sys.statements_with_full_table_scans", "Lọc lineitem theo l_partkey, đo bằng EXPLAIN, pg_stat_statements và pg_stat_user_tables"],
    ["Truy vấn con tương quan", "So lương với trung bình phòng trên emp_small, EXPLAIN có DEPENDENT SUBQUERY", "So giá trị đơn với trung bình của khách hàng trên orders, EXPLAIN có SubPlan"],
    ["Tranh chấp khóa", "Phiên A giữ khóa, phiên B chờ, quan sát sys.innodb_lock_waits", "Phiên A giữ khóa, phiên B chờ có lock_timeout, quan sát pg_stat_activity và pg_blocking_pids"],
    ["Mini λ-Tune", "SET GLOBAL tham số InnoDB và CREATE INDEX", "SET ở mức phiên (work_mem, random_page_cost và một số tham số khác) và CREATE INDEX"],
], "Bảng 2. Chuyển các kịch bản của Nhóm 3 sang PostgreSQL", [3.2, 5.9, 5.9])
p("Ở phần Mini λ-Tune, tham số shared_buffers không được đưa vào danh sách cho phép vì cần khởi động lại máy chủ. Thay vào đó, các tham số được đặt ở mức phiên, nên không ảnh hưởng tới các phiên khác.")
H("4.3. Danh sách kịch bản cho giám sát thời gian thực", 2)
table(["Kịch bản", "Mô tả", "Cách xử lý mong đợi"], [
    ["cpu_high", "Lặp 50 lần truy vấn tổng hợp kiểu TPC-H Q1 trên lineitem", "Hủy query, đặt statement_timeout"],
    ["missing_index", "Tìm lineitem theo l_partkey 10000 lần, cột không có index", "Tạo index trên l_partkey"],
    ["correlated_subquery", "Truy vấn con tương quan trên orders, chạy khoảng 30 đến 60 giây", "Viết lại bằng JOIN, hoặc tạo index trên o_custkey"],
    ["lock_holder", "Cập nhật đơn hàng o_orderkey = 1 rồi chờ 120 giây mới commit", "Tìm và kết thúc phiên giữ khóa"],
    ["lock_waiter", "Cập nhật cùng đơn hàng, bị chặn", "Đặt lock_timeout"],
    ["idle_in_transaction", "Mở transaction 90 giây không commit", "Đặt idle_in_transaction_session_timeout"],
    ["table_bloat", "Tắt autovacuum, cập nhật toàn bảng lineitem 3 lần", "VACUUM và bật lại autovacuum"],
    ["fix_vacuum", "Dọn dẹp sau table_bloat, không phải sự cố", "Không áp dụng"],
], "Bảng 3. Các kịch bản sự cố có sẵn", [3.8, 6.4, 4.8])
p("Người dùng có thể thêm kịch bản bằng cách tải lên một file SQL. Kịch bản chạy bằng user scenario_runner, chỉ có quyền trên schema tpch, và mỗi câu lệnh bị giới hạn 10 phút.")

H("5. Giao diện Streamlit", 1)
p("Giao diện có bảy tab. Tab 0 kiểm tra kết nối, dữ liệu và nguồn LLM. Tab 1, 2 và 3 giữ đúng luồng trình diễn của Nhóm 3 cho ba kịch bản. Người dùng bấm từng bước: gây sự cố, thu thập số đo, hỏi AI bằng prompt A và prompt B, xem điểm chấm tự động, sửa rồi đo lại và xem biểu đồ trước và sau. Tab 4 là Mini λ-Tune. Tab 5 là giám sát thời gian thực. Tab 6 tổng kết bảng chấm điểm và số đo.")
p("Ở tab 4, AI nhận một prompt theo mẫu Listing 1 của λ-Tune và trả về các câu lệnh SQL. Hệ thống chỉ giữ các lệnh khớp danh sách cho phép, các lệnh khác bị chặn và hiện cảnh báo. Mỗi cấu hình được áp dụng trên một phiên mới, chạy ba truy vấn có statement_timeout, rồi xóa các index vừa tạo. Trong lúc đo, monitor được tạm dừng để không báo nhầm.")

H("6. Luồng giám sát và phân tích tự động", 1)
p("Tab 5 chạy hoàn toàn tự động theo các bước sau.")
for i, t in enumerate([
    "Người dùng bấm Chạy một kịch bản, hoặc có một sự cố thật trên cơ sở dữ liệu.",
    "Monitor quét pg_stat_activity, pg_stat_statements và pg_stat_user_tables mỗi 5 giây. Khi vượt ngưỡng, monitor ghi một alert mới.",
    "Worker kiểm tra alert mới. Nếu công tắc Tự động phân tích đang bật, alert đã có ít nhất 20 giây và dấu hiệu vẫn còn ở lần quét gần nhất, worker đưa alert vào hàng đợi. Thời gian chờ này giúp bỏ qua các sự cố chỉ thoáng qua.",
    "Worker lấy từng alert trong hàng đợi, mỗi lần một alert, gọi công cụ chỉ đọc để lấy thêm số liệu, tra tri thức, rồi viết báo cáo.",
    "Trang Streamlit tự làm mới mỗi 3 giây. Trạng thái phân tích đi từ chờ phân tích sang đang phân tích rồi đã có báo cáo.",
    "Người dùng chọn một alert để xem tóm tắt, các bước suy luận có thể mở ra (công cụ, quan sát, kết luận), nguyên nhân gốc và phương án xử lý kèm câu SQL.",
    "Khi sự cố hết, sau 3 lần quét liên tiếp không còn dấu hiệu, alert tự chuyển sang đã giải quyết. Báo cáo vẫn được giữ lại.",
], 1):
    p(f"Bước {i}. {t}", WD_ALIGN_PARAGRAPH.LEFT)
p("Khi tắt công tắc Tự động phân tích, người dùng tự bấm nút Phân tích. Nút Phân tích lại cho phép chọn nguồn LLM khác cho cùng một alert. Công cụ của worker dùng kết nối chỉ đọc, nên AI chỉ đề xuất mà không tự sửa cơ sở dữ liệu.")
table(["Loại alert", "Điều kiện phát hiện"], [
    ["long_query", "Phiên đang chạy một câu lệnh quá 5 giây, không phải chờ khóa"],
    ["lock_wait", "Phiên chờ khóa quá 3 giây"],
    ["idle_in_transaction", "Transaction mở quá 10 giây mà không làm việc"],
    ["missing_index", "Một câu SELECT có điều kiện lọc được gọi từ 20 lần trong một chu kỳ, mỗi lần đọc từ 1000 block"],
    ["table_bloat", "Bảng có từ 100 nghìn dòng chết và số dòng chết từ 20% số dòng sống"],
    ["too_many_connections", "Số kết nối từ 80% max_connections"],
], "Bảng 4. Các loại alert và điều kiện phát hiện", [4, 11])

H("7. Nguồn LLM và cơ chế dự phòng", 1)
p("Hệ thống hỗ trợ năm nguồn: Claude, Gemini, OpenAI, Ollama và chế độ mock. API key đọc từ file .env và không được lưu trong repo. Model mặc định được kiểm tra trên tài liệu của nhà cung cấp ngày 15/10/2026: claude-sonnet-5-5 cho Claude và gemini-3.8-flash cho Gemini. Với Ollama, model mặc định là qwen2.5:7b, gọi tới địa chỉ host.docker.internal:11434 trên máy người dùng.")
p("Nếu nguồn đã chọn không có API key, hoặc gọi lỗi sau vài lần thử lại khi máy chủ quá tải, hệ thống tự chuyển sang Ollama và ghi chú việc chuyển này vào báo cáo. Chế độ mock không gọi LLM. Ở chế độ này, worker vẫn gọi công cụ thật trên cơ sở dữ liệu nhưng kết luận theo luật viết sẵn, dùng để kiểm thử luồng.")

H("8. Hướng dẫn chạy", 1)
p("Máy cần có Docker. Nếu muốn dùng LLM trên máy thì cài thêm Ollama và tải model bằng lệnh ollama pull qwen2.5:7b. Sau đó tải mã nguồn từ branch feature/realtime-streamlit-pg, sao chép file .env.example thành .env, điền API key nếu có, rồi chạy docker compose up --build -d. Lần đầu mất vài phút để build và sinh dữ liệu. Giao diện mở tại địa chỉ http://localhost:8501. File README trong repo có hướng dẫn chi tiết, gồm cách thêm kịch bản và cách xử lý lỗi thường gặp.")

H("9. Kết quả kiểm thử", 1)
p("Máy kiểm thử không có Docker, không có Ollama và không có API key. Vì vậy tác giả kiểm thử như sau: dùng PostgreSQL 17 cài trực tiếp, sinh dữ liệu bằng chính script seed.sh, chạy monitor, backend, worker và Streamlit bằng Python, và dùng chế độ mock cho phần AI. Sinh và nạp dữ liệu mất khoảng 8,6 giây, bảng lineitem có 600.572 dòng.")
p("Bảng 5 ghi lại kết quả của luồng tự động. Một script gọi API để chạy từng kịch bản, chờ báo cáo, dừng kịch bản rồi chờ alert chuyển sang đã giải quyết. Thời gian tính từ lúc bấm Chạy.")
table(["Kịch bản", "Alert đầu tiên (giây)", "Có báo cáo (giây)", "Dừng lúc (giây)", "Đã giải quyết (giây)"], [
    ["cpu_high", "9,0", "30,1", "30,1", "44,2"],
    ["missing_index", "10,0", "31,1", "31,2", "50,2"],
    ["correlated_subquery", "10,0", "30,2", "30,2", "50,3"],
    ["lock_holder và lock_waiter", "10,0", "31,2", "31,2", "45,3"],
    ["idle_in_transaction", "15,1", "35,2", "35,2", "49,3"],
    ["table_bloat và fix_vacuum", "5,0", "26,1", "26,2", "40,2"],
], "Bảng 5. Kết quả luồng tự động ở chế độ mock", [4.6, 2.6, 2.6, 2.6, 2.6])
p("Ở cả sáu lần chạy, báo cáo xuất hiện khoảng 20 giây sau alert đầu tiên, đúng với thời gian chờ đã đặt. Alert chuyển sang đã giải quyết sau khoảng 14 đến 20 giây kể từ khi dừng, tương ứng 3 lần quét. Kết luận ở chế độ mock đúng với từng loại sự cố. Ví dụ với correlated_subquery, worker thấy SubPlan trong kế hoạch thực thi và đề xuất viết lại bằng JOIN. Với missing_index, worker chỉ ra cột l_partkey không có index.")
p("Bảng 6 ghi lại kết quả các tab trình diễn từng bước, chạy bằng công cụ kiểm thử AppTest của Streamlit.")
table(["Nội dung", "Kết quả quan sát"], [
    ["Tab 1, thiếu index", "50 truy vấn mất 2,096 giây khi chưa có index và 0,013 giây sau khi tạo index"],
    ["Tab 2, truy vấn con", "Bản chậm mất 7,875 giây, EXPLAIN có SubPlan. Bản viết lại mất 0,063 giây, cùng kết quả 951 đơn"],
    ["Tab 3, tranh chấp khóa", "Công cụ thấy một phiên chờ, chỉ đúng pid phiên chặn. Sau khi phiên A ROLLBACK, phiên B chạy xong sau 2,9 giây chờ"],
    ["Chấm điểm (mock)", "Câu trả lời cho prompt B được chấm đúng nguyên nhân và có lệnh cụ thể ở cả ba kịch bản. Câu trả lời chung chung cho prompt A bị chấm sai"],
    ["Tab 4, Mini λ-Tune", "C0 mặc định 0,167 giây, C1 0,168 giây, C2 0,118 giây, C3 0,251 giây, cấu hình do mock đề xuất 0,112 giây. Lệnh ngoài danh sách cho phép bị chặn"],
], "Bảng 6. Kết quả các tab trình diễn từng bước", [4, 11])
p("Ngoài ra, tác giả kiểm tra thêm các điểm sau. Khi tắt Tự động phân tích, alert vẫn ở trạng thái chờ phân tích sau 35 giây. Khi bấm Phân tích lại với Claude mà không có API key, hệ thống ghi chú việc chuyển sang Ollama, rồi báo lỗi kết nối vì máy không có Ollama. Với chế độ mock, phân tích lại cho báo cáo ba bước. Tải lên một file SQL hợp lệ thì chạy được, tên kịch bản sai quy tắc bị từ chối. Bảy kiểm thử đơn vị cho chấm điểm, danh sách cho phép và cơ chế dự phòng đều đạt.")
p("Lưu ý về Bảng 6: chấm điểm ở chế độ mock chỉ kiểm tra luồng, không phản ánh chất lượng của LLM thật. Workload của Mini λ-Tune khá nhẹ trên SF 0.1 nên chênh lệch giữa các cấu hình nhỏ.")
p("Các phần chưa được kiểm thử gồm: chạy bằng docker compose và các Dockerfile; gọi LLM thật qua Claude, Gemini, OpenAI hoặc Ollama; thao tác bấm nút trên trình duyệt thật; và chạy trên PostgreSQL 16 như trong Docker, vì máy kiểm thử dùng PostgreSQL 17.")

H("10. Hạn chế", 1)
p("Luật phát hiện còn đơn giản và dựa trên ngưỡng cố định, nên đôi khi một sự cố sinh thêm alert phụ. Ví dụ kịch bản missing_index chạy trong một khối DO dài nên cũng tạo alert query chạy lâu, và kịch bản giữ khóa cũng tạo alert transaction mở lâu. Monitor không đo trực tiếp CPU của máy mà suy ra từ việc phiên đang chạy không phải chờ. Tri thức chẩn đoán còn ít và tìm theo từ khóa. Mô hình 7B chạy trên CPU khá chậm và có thể trả JSON sai định dạng. Tab 4 xóa mọi index phụ trên các bảng lineitem, orders, customer và part trước và sau mỗi lần đo, nên index tạo ở tab 1 cũng bị xóa. Bản demo không có suy luận theo cây và không có nhiều agent chuyên môn như bài báo gốc.")

H("11. Kết luận", 1)
p("Bản demo đã kết hợp hai phần. Phần thứ nhất là luồng trình diễn từng bước của Nhóm 3, giúp thấy rõ AI chẩn đoán tốt hơn khi có thẻ kinh nghiệm và số đo thật, và vì sao cấu hình do AI đề xuất phải được đo lại. Phần thứ hai là giám sát thời gian thực, trong đó sự cố được phát hiện và phân tích tự động, rồi tự đánh dấu đã giải quyết. Toàn bộ chạy trên máy cá nhân bằng PostgreSQL và dữ liệu TPC-H. Hướng mở rộng gồm chạy với LLM thật để so sánh chất lượng báo cáo giữa các mô hình, thêm công cụ như gợi ý index bằng hypopg, và dùng workload nặng hơn cho Mini λ-Tune.")

H("Ghi công", 1)
p("Luồng trình diễn ba kịch bản, prompt A và prompt B, thẻ kinh nghiệm, cách chấm điểm tự động và ý tưởng Mini λ-Tune trong bản demo được chuyển từ dự án \"LLM As DBA\" của Nhóm 3, môn Hệ cơ sở dữ liệu tiên tiến. Tác giả xin cảm ơn Nhóm 3 đã cho phép sử dụng.")

H("Tài liệu tham khảo", 1)
for r in [
    "[1] X. Zhou, G. Li, Z. Sun và cộng sự, D-Bot: Database Diagnosis System using Large Language Models, arXiv:2312.01454, 2023.",
    "[2] V. Giannakouris, I. Trummer, λ-Tune: Harnessing Large Language Models for Automated Database System Tuning, SIGMOD 2025.",
    "[3] Nhóm 3, Demo LLM As DBA, môn Hệ cơ sở dữ liệu tiên tiến, 2026.",
    "[4] TsinghuaDatabaseGroup, DB-GPT, https://github.com/TsinghuaDatabaseGroup/DB-GPT.",
    "[5] Transaction Processing Performance Council, TPC-H Benchmark, https://www.tpc.org/tpch/.",
    "[6] G. Rahn, tpch-kit, https://github.com/gregrahn/tpch-kit.",
    "[7] PostgreSQL Global Development Group, tài liệu PostgreSQL 16.",
    "[8] CrewAI, https://docs.crewai.com. Streamlit, https://docs.streamlit.io.",
]:
    p(r, WD_ALIGN_PARAGRAPH.LEFT)

doc.save(os.path.join(HERE, "bao-cao-dbot-demo.docx"))
print("saved")
